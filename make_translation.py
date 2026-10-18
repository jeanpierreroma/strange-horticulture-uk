"""Робить translation.json для патчера з книги перекладу: {ID: {хеш EN: UK}}.
Англійський текст гри у файл не потрапляє, лише короткий хеш для вибору варіанта.
Використання: python make_translation.py ../plan/strange-horticulture-uk.xlsx translation.json"""
import sys, json, hashlib
from openpyxl import load_workbook


def key(en):
    return hashlib.sha1(" ".join((en or "").split()).encode("utf-8")).hexdigest()[:12]


tr = {}
for fn, n, id_, _, en, _, uk, *_ in load_workbook(sys.argv[1]).active.iter_rows(min_row=2, values_only=True):
    if uk: tr.setdefault(id_, {})[key(en)] = uk.replace("\r\n", "\n")
with open(sys.argv[2], "w", encoding="utf-8") as f:
    json.dump(tr, f, ensure_ascii=False, indent=0, sort_keys=True)
print(f"ID: {len(tr)}")
